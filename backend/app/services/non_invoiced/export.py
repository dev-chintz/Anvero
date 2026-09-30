"""The exports of the non-invoiced sales record (docs/NON_INVOICED_SALES.md, section 4e): CSV, Excel
and PDF of the rows a report lists, with the columns the owner chooses, in the order chosen.

By default what the accountant asked for (2026-09-30): a running number, the date, the buyer's
first and last name, and the amount; then the range's total gross. The CSV keeps the conventions of
the earlier report's (UTF-8 with a BOM so Excel reads Polish letters, a comma as the decimal
separator, every value guarded against being read as a formula); Excel gets real numbers and dates;
the PDF names the seller, the range and when it was made, for the archive, in a font with Polish
letters (DejaVu Sans, app/assets/fonts, with its licence).
"""

import csv
import io
from collections.abc import Callable
from dataclasses import dataclass
from datetime import UTC, date, datetime
from decimal import Decimal
from pathlib import Path

from app.core.order_number import format_order_number
from app.models.non_invoiced import LedgerKind
from app.services.non_invoiced.classifier import REASON_TEXT, Category, Reason
from app.services.non_invoiced.ledger import business_date
from app.services.non_invoiced.report import Report, ReportRow

FONTS = Path(__file__).resolve().parents[2] / "assets" / "fonts"

CATEGORY_TEXT = {
    Category.EXEMPT_MAIL_ORDER: "Sprzedaż wysyłkowa zwolniona z kasy (poz. 41)",
    Category.PRIVATE_INVOICED: "Osoba prywatna z fakturą",
    Category.BUSINESS: "Zakup firmy",
    Category.NEEDS_REGISTER: "Wymaga kasy fiskalnej",
    Category.TO_REVIEW: "Do decyzji",
    Category.NOT_A_SALE: "Nie jest sprzedażą",
}

KIND_TEXT = {LedgerKind.SALE: "sprzedaż", LedgerKind.CORRECTION: "korekta"}

# a leading =, +, @ or - opens a formula in Excel/LibreOffice/Sheets if the cell is not guarded
_FORMULA_PREFIXES = ("=", "+", "@", "-")


def _guarded(value: str) -> str:
    return f"'{value}" if value and value[0] in _FORMULA_PREFIXES else value


def _money(value: Decimal) -> str:
    return f"{value:.2f}".replace(".", ",")


def _date(value: date | None) -> str:
    return value.strftime("%d.%m.%Y") if value else ""


def _reason(row: ReportRow) -> str:
    try:
        return REASON_TEXT[Reason(row.entry.reason)]
    except (ValueError, KeyError):
        return row.entry.reason


def _label_position(columns: list["Column"]) -> int:
    """Where the total row says "Razem": the first column after Lp. that is not the amount."""
    return next((i for i, c in enumerate(columns) if i > 0 and c.key != "amount"), 0)


def _buyer_name(row: ReportRow) -> str:
    return " ".join(p for p in (row.entry.buyer_first_name, row.entry.buyer_last_name) if p)


def _buyer_address(row: ReportRow) -> str:
    entry = row.entry
    city = " ".join(p for p in (entry.buyer_postal_code, entry.buyer_city) if p)
    return ", ".join(p for p in (entry.buyer_street, city) if p)


@dataclass(frozen=True)
class Column:
    key: str
    label: str
    # the value as text (CSV, PDF); Excel takes `value` where it gives one, a number or a date
    text: Callable[[int, ReportRow], str]
    value: Callable[[int, ReportRow], object] | None = None
    # its share of a PDF page's width, and whether it is set to the right (amounts)
    width: float = 1.0
    right: bool = False
    # names or reaches a person: an export holding it is logged (docs/GDPR.md)
    personal: bool = False


COLUMNS: dict[str, Column] = {
    c.key: c
    for c in (
        Column("lp", "Lp.", lambda i, _r: str(i), lambda i, _r: i, width=0.5),
        Column("entry_date", "Data", lambda _i, r: _date(r.entry.entry_date), lambda _i, r: r.entry.entry_date, width=1.0),
        Column("buyer_name", "Imię i nazwisko", lambda _i, r: _buyer_name(r), width=2.2, personal=True),
        Column("amount", "Kwota", lambda _i, r: _money(r.entry.amount), lambda _i, r: r.entry.amount, width=1.1, right=True),
        Column("kind", "Rodzaj", lambda _i, r: KIND_TEXT[r.entry.kind], width=0.9),
        Column("order_number", "Numer zamówienia", lambda _i, r: format_order_number(r.entry.order_number), width=1.2),
        Column("order_external_id", "Numer u marketplace'u", lambda _i, r: r.entry.order_external_id, width=2.4),
        Column("source", "Kanał", lambda _i, r: r.entry.source.value.capitalize(), width=0.8),
        Column("buyer_address", "Adres kupującego", lambda _i, r: _buyer_address(r), width=2.8, personal=True),
        Column("payment_operator", "Operator płatności", lambda _i, r: r.entry.payment_operator or "", width=1.0),
        Column("payment_id", "Identyfikator płatności", lambda _i, r: r.entry.payment_id or "", width=2.2),
        Column(
            "payout_date",
            "Data wypłaty",
            lambda _i, r: _date(business_date(r.entry.payout_at)) if r.entry.payout_at else "",
            lambda _i, r: business_date(r.entry.payout_at) if r.entry.payout_at else None,
            width=1.0,
        ),
        Column("payout_id", "Numer wypłaty", lambda _i, r: r.entry.payout_id or "", width=1.6),
        Column("category", "Kategoria", lambda _i, r: CATEGORY_TEXT[r.category], width=2.4),
        Column("reason", "Powód", lambda _i, r: _reason(r), width=3.0),
    )
}

# what the accountant asked for (2026-09-30), in this order
DEFAULT_COLUMNS = ["lp", "entry_date", "buyer_name", "amount"]

FORMATS = {"csv": ("text/csv", "csv"), "xlsx": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"), "pdf": ("application/pdf", "pdf")}


def resolve_columns(chosen: list[str] | None) -> list[Column]:
    """The chosen columns in the order given, "lp" always first; the default set when nothing was
    chosen. Raises KeyError naming the first key it does not know."""
    keys = [key for key in (chosen or DEFAULT_COLUMNS) if key != "lp"]
    return [COLUMNS["lp"], *(COLUMNS[key] for key in keys)]


def file_name(report: Report, extension: str) -> str:
    return f"ewidencja-bezrachunkowa-{report.date_from}-{report.date_to}.{extension}"


def _title(_report: Report) -> str:
    return "Ewidencja sprzedaży wysyłkowej zwolnionej z kasy rejestrującej (poz. 41)"


def _period(report: Report) -> str:
    return f"Okres: {_date(report.date_from)} – {_date(report.date_to)}"


# --- CSV ----------------------------------------------------------------------------------------


def to_csv(report: Report, columns: list[Column]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow([column.label for column in columns])
    for index, row in enumerate(report.listed, start=1):
        writer.writerow([_guarded(column.text(index, row)) for column in columns])
    keys = [c.key for c in columns]
    if "amount" in keys:
        total = [""] * len(columns)
        total[_label_position(columns)] = "Razem"
        total[keys.index("amount")] = _money(report.total)
        writer.writerow(total)
    return b"\xef\xbb\xbf" + buffer.getvalue().encode("utf-8")


# --- Excel --------------------------------------------------------------------------------------


def to_xlsx(report: Report, columns: list[Column]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    book = Workbook()
    sheet = book.active
    sheet.title = "Ewidencja"
    sheet.append([_title(report)])
    sheet["A1"].font = Font(bold=True, size=12)
    sheet.append([_period(report)])
    sheet.append([])
    header = 4
    sheet.append([column.label for column in columns])
    for cell in sheet[header]:
        cell.font = Font(bold=True)
    for index, row in enumerate(report.listed, start=1):
        values = []
        for column in columns:
            value = column.value(index, row) if column.value else column.text(index, row)
            values.append(value)
        sheet.append(values)
    last = header + len(report.listed)
    keys = [c.key for c in columns]
    if "amount" in keys:
        total = [None] * len(columns)
        total[_label_position(columns)] = "Razem"
        total[keys.index("amount")] = report.total
        sheet.append(total)
        for cell in sheet[last + 1]:
            cell.font = Font(bold=True)
    for position, column in enumerate(columns, start=1):
        letter = get_column_letter(position)
        sheet.column_dimensions[letter].width = max(8, int(column.width * 12))
        for cells in sheet.iter_rows(min_row=header + 1, max_row=sheet.max_row, min_col=position, max_col=position):
            for cell in cells:
                if column.key == "amount":
                    cell.number_format = "#,##0.00"
                    cell.alignment = Alignment(horizontal="right")
                elif column.key in ("entry_date", "payout_date"):
                    cell.number_format = "DD.MM.YYYY"
    sheet.freeze_panes = sheet.cell(row=header + 1, column=1)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


# --- PDF ----------------------------------------------------------------------------------------


def to_pdf(report: Report, columns: list[Column], seller: list[str], made_at: datetime | None = None) -> bytes:
    """The table on A4, landscape when the columns need it, with the seller named at the top."""
    from fpdf import FPDF
    from fpdf.fonts import FontFace

    made_at = made_at or datetime.now(UTC)
    wide = sum(c.width for c in columns) > 7.5

    class _Pdf(FPDF):
        def footer(self) -> None:
            self.set_y(-12)
            self.set_font("DejaVu", size=7)
            self.cell(0, 6, f"Strona {self.page_no()} z {{nb}}", align="R")

    pdf = _Pdf(orientation="L" if wide else "P", unit="mm", format="A4")
    pdf.add_font("DejaVu", "", str(FONTS / "DejaVuSans.ttf"))
    pdf.add_font("DejaVu", "B", str(FONTS / "DejaVuSans-Bold.ttf"))
    pdf.set_margins(12, 12, 12)
    pdf.set_auto_page_break(auto=True, margin=16)
    pdf.alias_nb_pages()
    pdf.add_page()

    pdf.set_font("DejaVu", size=8)
    for line in seller:
        pdf.cell(0, 4, line, new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)
    pdf.set_font("DejaVu", "B", 11)
    pdf.multi_cell(0, 5.5, _title(report), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("DejaVu", size=8.5)
    pdf.cell(0, 4.5, _period(report), new_x="LMARGIN", new_y="NEXT")
    made = f"Sporządzono: {business_date(made_at).strftime('%d.%m.%Y')}"
    if report.handed_over is not None:
        made += f" · przekazano księgowej {business_date(report.handed_over.handed_over_at).strftime('%d.%m.%Y')}"
    pdf.cell(0, 4.5, made, new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)

    pdf.set_font("DejaVu", size=8)
    widths = [c.width for c in columns]
    align = tuple("RIGHT" if c.right else "LEFT" for c in columns)
    with pdf.table(
        col_widths=widths,
        text_align=align,
        line_height=4.6,
        headings_style=FontFace(emphasis="BOLD", fill_color=(243, 245, 247)),
        borders_layout="HORIZONTAL_LINES",
    ) as table:
        heading = table.row()
        for column in columns:
            heading.cell(column.label)
        for index, row in enumerate(report.listed, start=1):
            cells = table.row()
            for column in columns:
                cells.cell(column.text(index, row))
        keys = [c.key for c in columns]
        if "amount" in keys:
            cells = table.row()
            for position, column in enumerate(columns):
                if column.key == "amount":
                    text = _money(report.total)
                elif position == _label_position(columns):
                    text = "Razem"
                else:
                    text = ""
                cells.cell(text, style=FontFace(emphasis="BOLD"))
    return bytes(pdf.output())


def export(report: Report, fmt: str, columns: list[Column], seller: list[str]) -> bytes:
    if fmt == "csv":
        return to_csv(report, columns)
    if fmt == "xlsx":
        return to_xlsx(report, columns)
    if fmt == "pdf":
        return to_pdf(report, columns, seller)
    raise ValueError(f"Unknown format: {fmt}")
